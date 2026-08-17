"""Build the document-generation MCP server.

Tools: ``create_excel`` (.xlsx), ``create_word`` (.docx), ``create_pdf`` (.pdf). Each renders
the given data into a real file, stores it, and returns a browser-reachable **download URL**
(LibreChat does not surface MCP-returned binary resources as downloads, so we serve them over
a plain HTTP route and hand back a link). GLM presents the link to the user as a Markdown link.

Storage: files are written to disk under ``DOCGEN_DATA_DIR`` (a Docker volume) so links survive
container restarts, bounded by ``DOCGEN_MAX_FILES`` (oldest evicted first).

Delivery route: ``/files/{id}/{filename}`` — the filename is in the URL path so the browser
always saves the file with the correct extension. PDFs are served ``inline`` (viewable in the
browser); spreadsheets/documents are served as ``attachment`` downloads (browsers can't render
them).
"""

from __future__ import annotations

import io
import shutil
import uuid
from pathlib import Path
from typing import Any

from docx import Document
from fastmcp import FastMCP
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from starlette.requests import Request
from starlette.responses import JSONResponse, Response

from .config import settings

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"

_MIME_BY_EXT = {".xlsx": XLSX_MIME, ".docx": DOCX_MIME, ".pdf": PDF_MIME}

_DATA_DIR = Path(settings.docgen_data_dir)
_DATA_DIR.mkdir(parents=True, exist_ok=True)


def _store_file(filename: str, data: bytes, mime: str) -> str:
    """Persist the file under DATA_DIR/<id>/<filename> and return the id. Evicts oldest
    files once the count exceeds the configured cap."""
    file_id = uuid.uuid4().hex
    dest = _DATA_DIR / file_id
    dest.mkdir(parents=True, exist_ok=True)
    (dest / filename).write_bytes(data)

    dirs = [p for p in _DATA_DIR.iterdir() if p.is_dir()]
    if len(dirs) > settings.docgen_max_files:
        for old in sorted(dirs, key=lambda p: p.stat().st_mtime)[: -settings.docgen_max_files]:
            shutil.rmtree(old, ignore_errors=True)
    return file_id


def _deliver(filename: str, data: bytes, mime: str, ext: str) -> str:
    """Store the file and return the user-facing download instruction + link."""
    if not filename.lower().endswith(ext):
        filename = f"{filename}{ext}"
    file_id = _store_file(filename, data, mime)
    # Filename in the URL path guarantees the browser saves it with the right extension.
    url = f"{settings.docgen_public_base_url.rstrip('/')}/files/{file_id}/{filename}"
    return (
        f"File '{filename}' created.\n"
        f"Download URL: {url}\n"
        f"Present this to the user as a clickable Markdown link: [{filename}]({url})"
    )


def build_server() -> FastMCP:
    mcp = FastMCP(name="docgen")

    # Two routes to the same handler. The filename is in the URL path (…/{id}/{name}) so the
    # browser keeps the extension even if it ignores Content-Disposition. Lookup uses only
    # file_id; the {filename} segment is cosmetic.
    @mcp.custom_route("/files/{file_id}", methods=["GET"])
    @mcp.custom_route("/files/{file_id}/{filename}", methods=["GET"])
    async def download_file(request: Request) -> Response:
        file_id = request.path_params["file_id"]
        if not file_id.isalnum():  # ids are uuid4 hex; reject anything else (path-traversal guard)
            return JSONResponse({"error": "invalid file id"}, status_code=400)

        folder = _DATA_DIR / file_id
        files = [p for p in folder.iterdir() if p.is_file()] if folder.is_dir() else []
        if not files:
            return JSONResponse({"error": "file not found or expired"}, status_code=404)

        fp = files[0]
        mime = _MIME_BY_EXT.get(fp.suffix.lower(), "application/octet-stream")
        # PDFs are viewable in-browser (inline); spreadsheets/docs must download (attachment).
        disposition = "inline" if mime == PDF_MIME else "attachment"
        return Response(
            content=fp.read_bytes(),
            media_type=mime,
            headers={"Content-Disposition": f'{disposition}; filename="{fp.name}"'},
        )

    @mcp.tool
    def create_excel(
        filename: str,
        headers: list[str],
        rows: list[list[Any]],
        sheet_name: str = "Sheet1",
    ) -> str:
        """Generate an Excel (.xlsx) workbook from tabular data and return a **download link**.

        Use for exporting tabular/spreadsheet data. Always present the returned download URL
        to the user as a clickable Markdown link, e.g. ``[filename](url)``.

        Args:
            filename: desired file name, e.g. "chargers.xlsx".
            headers: column header names (written as a bold first row).
            rows: list of rows; each row is a list of cell values aligned to ``headers``.
            sheet_name: worksheet tab name (max 31 chars).
        """
        wb = Workbook()
        ws = wb.active
        ws.title = (sheet_name or "Sheet1")[:31]

        if headers:
            ws.append(list(headers))
            for cell in ws[1]:
                cell.font = Font(bold=True)
        for row in rows:
            ws.append(list(row))

        for col_idx, _ in enumerate(headers or (rows[0] if rows else []), start=1):
            width = 12
            for row in ([headers] if headers else []) + rows:
                if col_idx <= len(row) and row[col_idx - 1] is not None:
                    width = max(width, min(60, len(str(row[col_idx - 1])) + 2))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

        buf = io.BytesIO()
        wb.save(buf)
        return _deliver(filename, buf.getvalue(), XLSX_MIME, ".xlsx")

    @mcp.tool
    def create_word(
        filename: str,
        title: str = "",
        paragraphs: list[str] | None = None,
        table_headers: list[str] | None = None,
        table_rows: list[list[Any]] | None = None,
    ) -> str:
        """Generate a Word (.docx) document and return a **download link**.

        Use for narrative reports/summaries. Combine a title, body paragraphs, and an optional
        table. Always present the returned download URL as a clickable Markdown link.

        Args:
            filename: desired file name, e.g. "fleet_report.docx".
            title: document title (rendered as the top heading).
            paragraphs: body paragraphs, in order.
            table_headers: optional column headers for a table after the paragraphs.
            table_rows: optional table rows (list of lists) aligned to ``table_headers``.
        """
        if not (title or paragraphs or table_headers):
            title = filename  # guard: never emit an empty document if the model omitted content
        doc = Document()
        if title:
            doc.add_heading(title, level=0)
        for para in paragraphs or []:
            doc.add_paragraph(str(para))

        if table_headers:
            table = doc.add_table(rows=1, cols=len(table_headers))
            table.style = "Table Grid"
            for i, h in enumerate(table_headers):
                table.rows[0].cells[i].text = str(h)
            for row in table_rows or []:
                cells = table.add_row().cells
                for i, val in enumerate(row):
                    if i < len(cells):
                        cells[i].text = "" if val is None else str(val)

        buf = io.BytesIO()
        doc.save(buf)
        return _deliver(filename, buf.getvalue(), DOCX_MIME, ".docx")

    @mcp.tool
    def create_pdf(
        filename: str,
        title: str = "",
        paragraphs: list[str] | None = None,
        table_headers: list[str] | None = None,
        table_rows: list[list[Any]] | None = None,
    ) -> str:
        """Generate a PDF (.pdf) document and return a **download link**.

        Use for printable reports/summaries. Combine a title, body paragraphs, and an optional
        table. Always present the returned download URL as a clickable Markdown link.

        Args:
            filename: desired file name, e.g. "fleet_summary.pdf".
            title: document title (rendered as the top heading).
            paragraphs: body paragraphs, in order.
            table_headers: optional column headers for a table after the paragraphs.
            table_rows: optional table rows (list of lists) aligned to ``table_headers``.
        """
        if not (title or paragraphs or table_headers):
            title = filename  # guard: never emit a 0-page PDF if the model omitted content
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4, title=title or filename)
        styles = getSampleStyleSheet()
        story: list[Any] = []

        if title:
            story.append(Paragraph(title, styles["Title"]))
            story.append(Spacer(1, 0.4 * cm))
        for para in paragraphs or []:
            story.append(Paragraph(str(para), styles["BodyText"]))
            story.append(Spacer(1, 0.2 * cm))

        if table_headers:
            data = [[str(h) for h in table_headers]]
            for row in table_rows or []:
                data.append(["" if v is None else str(v) for v in row])
            table = Table(data, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563eb")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("PADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.append(Spacer(1, 0.3 * cm))
            story.append(table)

        doc.build(story)
        return _deliver(filename, buf.getvalue(), PDF_MIME, ".pdf")

    return mcp
