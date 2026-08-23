from pathlib import Path
from kb_mcp.parse import parse_file, SUPPORTED


def test_parses_markdown(tmp_path: Path):
    f = tmp_path / "doc.md"
    f.write_text("# Title\n\nBody paragraph about OCPP error 42.", encoding="utf-8")
    segs = parse_file(f)
    assert segs and any("OCPP error 42" in s["text"] for s in segs)
    assert all("location" in s for s in segs)


def test_parses_txt_and_html(tmp_path: Path):
    (tmp_path / "a.txt").write_text("plain text body", encoding="utf-8")
    (tmp_path / "b.html").write_text("<h1>H</h1><p>html body</p>", encoding="utf-8")
    assert "plain text body" in parse_file(tmp_path / "a.txt")[0]["text"]
    assert "html body" in parse_file(tmp_path / "b.html")[0]["text"]


def test_unsupported_extension_returns_empty(tmp_path: Path):
    (tmp_path / "x.bin").write_bytes(b"\x00\x01")
    assert parse_file(tmp_path / "x.bin") == []
    assert ".pdf" in SUPPORTED and ".md" in SUPPORTED


def _make_docx(path, paragraphs, table_rows=None, heading=None):
    from docx import Document
    doc = Document()
    if heading:
        doc.add_heading(heading, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        t = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    doc.save(path)


def test_parses_docx_paragraphs_and_tables(tmp_path: Path):
    f = tmp_path / "runbook.docx"
    _make_docx(f, ["Connector lock failure recovery."],
               table_rows=[["Code", "Meaning"], ["E42", "comms lost"]],
               heading="Stuck Connector")
    segs = parse_file(f)
    body = " ".join(s["text"] for s in segs)
    assert "Connector lock failure recovery." in body
    assert "E42" in body and "comms lost" in body
    assert any(s["location"] == "Stuck Connector" for s in segs)


def test_empty_docx_yields_no_segments(tmp_path: Path):
    f = tmp_path / "empty.docx"
    _make_docx(f, [])
    assert parse_file(f) == []


def test_docx_is_supported(tmp_path: Path):
    assert ".docx" in SUPPORTED
